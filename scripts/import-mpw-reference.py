"""Read the supplied historical workbook; write explicit, traceable summary records.

Usage: python scripts/import-mpw-reference.py path/to/overview.xlsx
Only reads the source workbook. Draft and cross-project average rows are excluded.
"""
import datetime
import json
import pathlib
import sys
import openpyxl

source = pathlib.Path(sys.argv[1])
book = openpyxl.load_workbook(source, data_only=True)
records = []

def merged_value(sheet, row, col):
    for area in sheet.merged_cells.ranges:
        if area.min_row <= row <= area.max_row and area.min_col <= col <= area.max_col:
            return sheet.cell(area.min_row, area.min_col).value
    return sheet.cell(row, col).value

def numeric(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)

def date(value):
    return value.date().isoformat() if isinstance(value, datetime.datetime) else ''

def add(sheet, row, col, platform, project, slot, block, metric, parameter, unit, measured, method, notes='', mode='Unknown'):
    value = sheet.cell(row, col).value
    if not numeric(value):
        return
    cell = sheet.cell(row, col).coordinate
    records.append(dict(id=f'reference:{sheet.title}:{cell}', project=f'MPW{int(project)}', platform=platform,
        slot=f'Slot{slot}', step='Unknown', buildingBlock=block, opticalMode=mode, metric=metric, parameter=parameter,
        unit=unit, date=date(measured), method=str(method or 'Unknown'), summaryValue=value, chips=[], source='reference',
        sourceRef=f'{source.name} / {sheet.title}!{cell}', notes=notes, referenceCell=cell))

for title, heater_only in [('MPW_RUN-prop-Loss', False), ('MPW_RUN-Heaters', True)]:
    sheet = book[title]
    current_platform = 'Unknown'
    for row in range(3, sheet.max_row + 1):
        platform_label = merged_value(sheet, row, 1)
        if platform_label:
            current_platform = str(platform_label).strip()
        project = merged_value(sheet, row, 2)
        slot = sheet.cell(row, 3).value
        if not numeric(project) or slot is None or 'based' in str(slot):
            continue
        platform = current_platform
        # The active platform label is visually split across two rows in the source.
        if (not heater_only and 21 <= row <= 26) or (heater_only and 19 <= row <= 24):
            platform = 'SOI: 220nm Active'
        if platform == 'Platform':
            platform = 'Unknown'
        if not heater_only:
            note = str(sheet.cell(row, 8).value or '')
            for col, geometry in [(4, 'Rib'), (5, 'Strip')]:
                add(sheet, row, col, platform, project, slot, f'{geometry} waveguide', 'propagation', 'Propagation loss', 'dB/cm', sheet.cell(row, 6).value, sheet.cell(row, 7).value, note, '1550nm (polarisation unspecified)')
        offset = 3 if heater_only else 8
        date_col = 12 if heater_only else 17
        method_col = date_col + 1
        for width, delta in [(2, 0), (5, 1)]:
            for start, parameter, unit in [(offset + 1, 'Ppi', 'mW/pi'), (offset + 3, 'Vpi', 'V'), (offset + 5, 'Vpi × L', 'V·um'), (offset + 7, 'FSR', 'unit unspecified')]:
                add(sheet, row, start + delta, platform, project, slot, f'MZI heater {width}um width / 200um length', 'heater', parameter, unit, sheet.cell(row, date_col).value, sheet.cell(row, method_col).value,
                    'Historical summary. FSR unit not stated in source; Vpi × L uses the stated 200um length. Separate source-sheet values retained when they disagree.')

sheet = book['MPW_RUN-MMIs (NEW)']
device = 'Unknown'
for row in range(4, sheet.max_row + 1):
    if sheet.cell(row, 1).value:
        device = str(sheet.cell(row, 1).value)
    project = merged_value(sheet, row, 2)
    slot = sheet.cell(row, 3).value
    if not numeric(project) or slot is None or 'based' in str(slot):
        continue
    for col, geometry in [(4, 'Rib'), (5, 'Strip')]:
        for delta, parameter in [(0, 'Insertion loss (raw)'), (2, 'Insertion loss (splitting corrected)'), (4, 'Insertion loss (absolute corrected)')]:
            add(sheet, row, col + delta, 'Platform unspecified in MMI sheet', project, slot, f'{geometry} {device}', 'insertion', parameter, 'dB/unit', sheet.cell(row, 10).value, sheet.cell(row, 11).value,
                'Raw, corrected and absolute-corrected columns are distinct conventions. No additional 3dB subtraction has been applied. Derived cross/bar averages are labelled by device.')

# Do not duplicate identical heater summaries copied between the two source sheets.
deduplicated = []
seen = set()
for record in records:
    key = tuple(str(record.get(k, '')) for k in ['platform', 'project', 'slot', 'buildingBlock', 'parameter', 'date', 'summaryValue'])
    if key in seen:
        continue
    seen.add(key)
    deduplicated.append(record)
target = pathlib.Path(__file__).resolve().parents[1] / 'public/mpw-database/reference-records.json'
target.parent.mkdir(parents=True, exist_ok=True)
target.write_text(json.dumps(deduplicated, indent=2, ensure_ascii=False) + '\n', encoding='utf-8')
print(json.dumps({'records': len(deduplicated), 'metrics': {m: sum(r['metric'] == m for r in deduplicated) for m in ['propagation', 'insertion', 'heater']}}))
